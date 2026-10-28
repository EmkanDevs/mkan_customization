import json
import copy
import traceback
from io import BytesIO

import frappe
from frappe import _
from frappe.desk.query_report import get_column_as_dict, run
from frappe.utils import cint, cstr, flt, get_datetime, strip_html
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

DATE_FIELDTYPES = {"Date"}
DATETIME_FIELDTYPES = {"Datetime"}


class _ReportFilters(frappe._dict):
    """frappe._dict that is always truthy, even when empty."""

    def __bool__(self):
        return True


def _to_report_filters(value):
    if isinstance(value, dict):
        return _ReportFilters({k: _to_report_filters(v) for k, v in value.items()})
    if isinstance(value, (list, tuple)):
        return [_to_report_filters(v) for v in value]
    return value


def normalize_report_filters(raw):
    if raw is None or raw == "":
        return _ReportFilters()
    if isinstance(raw, str):
        try:
            raw = json.loads(raw)
        except ValueError as e:
            frappe.throw(_("Report filters are not valid JSON: {0}").format(e))
        if raw is None:
            return _ReportFilters()
    if not isinstance(raw, dict):
        frappe.throw(_("Report filters must be an object, got {0}.").format(type(raw).__name__))
    return _to_report_filters(copy.deepcopy(dict(raw)))


# ---------------------------------------------------------------------------
# Shared workbook builder + download response
# ---------------------------------------------------------------------------
def _build_and_send(title, labels, fieldnames, fieldtypes, rows, filename):
    wb = Workbook()
    ws = wb.active
    ws.title = (title or "Sheet1")[:31]

    header_font = Font(bold=True, size=10)
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")

    for idx, label in enumerate(labels, start=1):
        cell = ws.cell(row=1, column=idx, value=label)
        cell.font = header_font
        cell.alignment = center
        cell.border = border
    ws.row_dimensions[1].height = 20

    max_width = [len(str(lbl)) + 4 for lbl in labels]

    for r_idx, row in enumerate(rows):
        if isinstance(row, dict):
            values = [row.get(fn, "") for fn in fieldnames]
        else:
            values = [row[i] if i < len(row) else "" for i in range(len(fieldnames))]

        excel_row = 2 + r_idx
        for c_idx, (value, ftype) in enumerate(zip(values, fieldtypes), start=1):
            if isinstance(value, bool):
                value = "Yes" if value else "No"
            if value is None:
                value = ""

            cell = ws.cell(row=excel_row, column=c_idx)

            if ftype in ("Currency", "Float") and isinstance(value, (int, float)):
                cell.value = flt(value, 2)
                cell.number_format = "#,##0.00"
                cell.alignment = right_align
            elif ftype == "Percent" and isinstance(value, (int, float)):
                cell.value = flt(value, 2) / 100
                cell.number_format = "0.00%"
                cell.alignment = right_align
            elif ftype == "Int" and isinstance(value, (int, float)):
                cell.value = cint(value)
                cell.number_format = "#,##0"
                cell.alignment = right_align
            elif ftype in DATE_FIELDTYPES and value:
                try:
                    cell.value = get_datetime(value).date()
                    cell.number_format = "dd-mm-yyyy"
                except Exception:
                    cell.value = cstr(value)
                cell.alignment = center
            elif ftype in DATETIME_FIELDTYPES and value:
                try:
                    cell.value = get_datetime(value)
                    cell.number_format = "dd-mm-yyyy hh:mm"
                except Exception:
                    cell.value = cstr(value)
                cell.alignment = center
            else:
                if isinstance(value, str) and "<" in value and ">" in value:
                    value = strip_html(value)
                if isinstance(value, str) and value.startswith("="):
                    value = f"'{value}"
                cell.value = value
                cell.alignment = left_align

            cell.border = border

            text_len = len(cstr(value)) + 2
            if text_len > max_width[c_idx - 1]:
                max_width[c_idx - 1] = min(text_len, 45)

    for idx, width in enumerate(max_width, 1):
        ws.column_dimensions[get_column_letter(idx)].width = min(max(width, 10), 45)

    ws.freeze_panes = "A2"

    buffer = BytesIO()
    wb.save(buffer)

    frappe.local.response.filename = filename
    frappe.local.response.filecontent = buffer.getvalue()
    frappe.local.response.type = "download"
    frappe.local.response.content_type = (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


# ---------------------------------------------------------------------------
# 1. Query / Script Report export
# ---------------------------------------------------------------------------
@frappe.whitelist()
def export_query_report(form_params):
    if isinstance(form_params, str):
        form_params = json.loads(form_params)
    if not isinstance(form_params, dict):
        frappe.throw(_("Invalid parameters."))

    form_params = frappe._dict(form_params)
    report_name = form_params.get("report_name")
    if not report_name:
        frappe.throw(_("report_name is required."))

    filters = normalize_report_filters(form_params.get("filters"))

    # Normalize date_range
    if "date_range" in filters and isinstance(filters["date_range"], str):
        try:
            filters["date_range"] = json.loads(filters["date_range"])
        except ValueError:
            if "," in filters["date_range"]:
                filters["date_range"] = [d.strip() for d in filters["date_range"].split(",")]
            elif " to " in filters["date_range"]:
                filters["date_range"] = [d.strip() for d in filters["date_range"].split(" to ")]

    if not filters.get("date_range") and filters.get("from_date") and filters.get("to_date"):
        filters["date_range"] = [filters.get("from_date"), filters.get("to_date")]

    try:
        data = run(report_name, filters, ignore_prepared_report=True, are_default_filters=False)
    except (frappe.PermissionError, frappe.ValidationError):
        raise
    except Exception as e:
        frappe.log_error(message=traceback.format_exc(), title=f"Excel export failed: {report_name}")
        frappe.throw(
            _("Could not export report {0}: {1}: {2}").format(report_name, type(e).__name__, e),
            title=_("Excel export failed"),
        )
    data = frappe._dict(data)

    if "columns" not in data:
        frappe.throw(
            _("Report {0} returned no columns (keys: {1}).").format(report_name, ", ".join(data.keys()))
        )

    col_dicts = [get_column_as_dict(col) for col in (data.get("columns") or [])]
    fieldnames = [c.get("fieldname") for c in col_dicts]
    labels = [c.get("label") or c.get("fieldname") for c in col_dicts]
    fieldtypes = [c.get("fieldtype") or "Data" for c in col_dicts]

    _build_and_send(
        report_name,
        labels,
        fieldnames,
        fieldtypes,
        data.get("result") or [],
        frappe.scrub(report_name) + ".xlsx",
    )


# ---------------------------------------------------------------------------
# 2. Report View export (current filters, sort, visible columns)
# ---------------------------------------------------------------------------
@frappe.whitelist()
def export_list_view(form_params):
    if isinstance(form_params, str):
        form_params = json.loads(form_params)
    form_params = frappe._dict(form_params or {})

    doctype = form_params.get("doctype")
    if not doctype:
        frappe.throw(_("doctype is required."))

    frappe.has_permission(doctype, "export", throw=True)

    meta = frappe.get_meta(doctype)
    filters = form_params.get("filters") or []
    or_filters = form_params.get("or_filters") or []

    requested = form_params.get("columns") or []
    valid_standard = {"name", "owner", "creation", "modified", "modified_by", "docstatus", "idx"}
    columns = []
    for c in requested:
        fn = c.get("fieldname")
        if not fn or any(x["fieldname"] == fn for x in columns):
            continue
        df = meta.get_field(fn)
        if df:
            if df.fieldtype in frappe.model.no_value_fields:
                continue
            columns.append({"fieldname": fn, "label": df.label or fn, "fieldtype": df.fieldtype})
        elif fn in valid_standard:
            ftype = {
                "creation": "Datetime",
                "modified": "Datetime",
                "docstatus": "Int",
                "idx": "Int",
            }.get(fn, "Data")
            columns.append(
                {
                    "fieldname": fn,
                    "label": c.get("label") or fn.replace("_", " ").title(),
                    "fieldtype": ftype,
                }
            )

    if not columns:
        columns = [{"fieldname": "name", "label": _("ID"), "fieldtype": "Data"}]

    sort_by = form_params.get("sort_by") or "modified"
    sort_order = (form_params.get("sort_order") or "desc").lower()
    if sort_order not in ("asc", "desc"):
        sort_order = "desc"
    if sort_by not in valid_standard and not meta.get_field(sort_by):
        sort_by = "modified"
    order_by = f"`tab{doctype}`.`{sort_by}` {sort_order}"

    rows = frappe.get_list(
        doctype,
        fields=[c["fieldname"] for c in columns],
        filters=filters,
        or_filters=or_filters,
        order_by=order_by,
        limit_page_length=0,
    )

    _build_and_send(
        doctype,
        [c["label"] for c in columns],
        [c["fieldname"] for c in columns],
        [c["fieldtype"] for c in columns],
        rows,
        frappe.scrub(doctype) + ".xlsx",
    )